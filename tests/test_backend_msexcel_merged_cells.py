from collections.abc import Iterator
from datetime import date
from pathlib import Path

from docling_core.transforms.serializer.markdown import MarkdownParams
from docling_core.transforms.serializer.markdown_excel import (
    MsExcelMarkdownDocSerializer,
)
from docling_core.types.doc import ContentLayer
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.worksheet.merge import MergedCellRange

from docling.backend.msexcel_backend import MsExcelDocumentBackend
from docling.datamodel.base_models import InputFormat
from docling.datamodel.document import DoclingDocument, InputDocument
from docling.document_converter import DocumentConverter


class _TrackingMergedRanges(set[MergedCellRange]):
    def __init__(self, ranges: set[MergedCellRange]) -> None:
        super().__init__(ranges)
        self.iteration_count = 0

    def __iter__(self) -> Iterator[MergedCellRange]:
        self.iteration_count += 1
        return super().__iter__()


def _convert_workbook(workbook: Workbook, file_path: Path) -> DoclingDocument:
    workbook.save(file_path)
    converter = DocumentConverter(allowed_formats=[InputFormat.XLSX])
    return converter.convert(file_path).document


def test_merged_cells_preserve_table_semantics(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "merged"

    sheet["A1"] = "Header"
    sheet.merge_cells("A1:C1")
    sheet["D1"] = "Plain"

    sheet["A2"] = 42
    sheet.merge_cells("A2:A4")
    sheet["B2"] = date(2026, 7, 17)
    sheet.merge_cells("B2:D3")
    sheet["B4"] = "Footer"
    sheet["C4"] = "C"
    sheet["D4"] = "D"
    sheet["D4"].comment = Comment("Synthetic note", "Codex")

    sheet.merge_cells("F1:G2")
    sheet["H1"] = "After empty merge"

    doc = _convert_workbook(workbook, tmp_path / "merged-semantics.xlsx")

    assert len(doc.tables) == 2
    main_table = next(
        table
        for table in doc.tables
        if any(cell.text == "Header" for cell in table.data.table_cells)
    )
    empty_anchor_table = next(table for table in doc.tables if table != main_table)

    assert main_table.prov[0].bbox.as_tuple() == (0.0, 0.0, 4.0, 4.0)
    assert main_table.data.num_rows == 4
    assert main_table.data.num_cols == 4

    cells_by_text = {cell.text: cell for cell in main_table.data.table_cells}
    assert (cells_by_text["Header"].row_span, cells_by_text["Header"].col_span) == (
        1,
        3,
    )
    assert (cells_by_text["42"].row_span, cells_by_text["42"].col_span) == (3, 1)
    date_cell = next(
        cell
        for cell in main_table.data.table_cells
        if cell.text.startswith("2026-07-17")
    )
    assert (date_cell.row_span, date_cell.col_span) == (2, 3)
    assert len(main_table.data.table_cells) == 7

    assert empty_anchor_table.prov[0].bbox.as_tuple() == (5.0, 0.0, 8.0, 2.0)
    assert empty_anchor_table.data.num_rows == 2
    assert empty_anchor_table.data.num_cols == 3
    empty_anchor = next(
        cell
        for cell in empty_anchor_table.data.table_cells
        if cell.start_row_offset_idx == 0 and cell.start_col_offset_idx == 0
    )
    assert empty_anchor.text == ""
    assert (empty_anchor.row_span, empty_anchor.col_span) == (2, 2)
    assert any(
        text.content_layer == ContentLayer.NOTES and "Synthetic note" in text.text
        for text in doc.texts
    )

    markdown = (
        MsExcelMarkdownDocSerializer(
            doc=doc,
            params=MarkdownParams(compact_tables=True),
        )
        .serialize()
        .text
    )
    assert "Header" in markdown
    assert "After empty merge" in markdown


def test_empty_and_unmerged_sheets_do_not_regress(tmp_path: Path) -> None:
    workbook = Workbook()
    workbook.active.title = "empty"
    sheet = workbook.create_sheet("unmerged")
    sheet["A1"] = "left"
    sheet["B1"] = "right"

    doc = _convert_workbook(workbook, tmp_path / "empty-and-unmerged.xlsx")

    assert len(doc.pages) == 2
    assert doc.pages[1].size.as_tuple() == (0.0, 0.0)
    assert doc.pages[2].size.as_tuple() == (2.0, 1.0)
    assert len(doc.tables) == 1
    assert [cell.text for cell in doc.tables[0].data.table_cells] == ["left", "right"]


def test_merged_ranges_are_indexed_once_per_worksheet(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    for row in range(1, 51):
        sheet.cell(row=row, column=1, value=f"row-{row}")
        sheet.merge_cells(
            start_row=row,
            start_column=1,
            end_row=row,
            end_column=3,
        )
        sheet.cell(row=row, column=4, value=row)

    file_path = tmp_path / "tracked-merged-ranges.xlsx"
    workbook.save(file_path)
    in_doc = InputDocument(
        path_or_stream=file_path,
        format=InputFormat.XLSX,
        filename=file_path.stem,
        backend=MsExcelDocumentBackend,
    )
    backend = MsExcelDocumentBackend(in_doc=in_doc, path_or_stream=file_path)
    loaded_sheet = backend.workbook.active
    tracked_ranges = _TrackingMergedRanges(loaded_sheet.merged_cells.ranges)
    loaded_sheet.merged_cells.ranges = tracked_ranges

    tables, _ = backend._find_data_tables(loaded_sheet)

    assert len(tables) == 1
    assert tracked_ranges.iteration_count == 1
